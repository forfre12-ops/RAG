"""부서×정보유형 카테고리 통계·조회 + 검수 결과 엑셀 내보내기 (관리자 전용).

이상민(영업비밀보호센터 주임) 요청 #1(카테고리별 데이터 분류조회)·#2(가이드별 생성 건수
통계)·#4(의견 입력+출력) 대응. 의견은 새 입력칸을 안 만들고 결정의 사유(reason)를 그대로
쓴다(2026-10-02 사용자 결정 A안).
"""

from __future__ import annotations

import io
import json
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from koipa.api._jwt_auth import JWTClaims, require_auth
from koipa.api.app import app
from koipa.services import proxy_gold_candidate_service as pgs

API = "/api/v1"
ADMIN = {"mode": "jwt", "claims": JWTClaims(sub="admin-kim", roles=("admin",), exp=9999999999),
         "actor_role": "admin", "actor_roles": ["admin"]}
REVIEWER = {"mode": "jwt", "claims": JWTClaims(sub="expert-a", roles=("reviewer",), exp=9999999999),
            "actor_role": "reviewer", "actor_roles": ["reviewer"]}


def _write_candidate(root: Path, doc_id: str, *, department: str, info_type: str, grade: str) -> None:
    root.mkdir(parents=True, exist_ok=True)
    (root / f"{doc_id}_review.md").write_text(f"본문 {doc_id}", encoding="utf-8")
    meta = {
        "doc_id": doc_id, "intended_label": grade, "document_origin": "synthetic",
        "department": department, "info_type": info_type, "candidate_status": "proposed",
    }
    (root / f"{doc_id}.metadata.json").write_text(json.dumps(meta, ensure_ascii=False), encoding="utf-8")


@pytest.fixture
def pool(tmp_path, monkeypatch):
    root = tmp_path / "pool"
    monkeypatch.setattr(pgs, "_DEFAULT_ROOT", root)
    _write_candidate(root, "D1", department="생산제조", info_type="생산기술 정보", grade="S1")
    _write_candidate(root, "D2", department="생산제조", info_type="생산기술 정보", grade="S1")
    _write_candidate(root, "D3", department="경영", info_type="경영실적 정보", grade="S2")
    return root


@pytest.fixture
def client():
    return TestClient(app)


def _as(who: dict) -> None:
    app.dependency_overrides[require_auth] = lambda: who


def teardown_function(_fn):
    app.dependency_overrides.pop(require_auth, None)


# ── category-stats ───────────────────────────────────────────────────────


def test_category_stats_groups_by_department_info_type_grade(pool, client):
    _as(ADMIN)
    r = client.get(f"{API}/golden/candidates/category-stats")
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["total"] == 3
    assert {"생산제조", "경영"} == set(body["departments"])
    counts = {(row["department"], row["info_type"], row["grade"]): row["count"] for row in body["rows"]}
    assert counts[("생산제조", "생산기술 정보", "S1")] == 2
    assert counts[("경영", "경영실적 정보", "S2")] == 1


def test_category_stats_is_admin_only(pool, client):
    _as(REVIEWER)
    r = client.get(f"{API}/golden/candidates/category-stats")
    assert r.status_code == 403


# ── export ────────────────────────────────────────────────────────────────


def test_export_with_filter_returns_all_matching_rows_as_xlsx(pool, client):
    _as(ADMIN)
    r = client.post(f"{API}/golden/candidates/export", json={"department": "생산제조"})
    assert r.status_code == 200, r.text
    assert r.headers["content-type"].startswith("application/vnd.openxmlformats")
    wb = load_workbook(io.BytesIO(r.content))
    ws = wb.active
    doc_ids = [row[0].value for row in ws.iter_rows(min_row=2)]
    assert sorted(doc_ids) == ["D1", "D2"]


def test_export_with_doc_ids_returns_exactly_those_and_ignores_filters(pool, client):
    """doc_ids 를 주면 그게 이긴다 — 필터를 같이 보내도 무시된다(선택 내보내기가 우선)."""
    _as(ADMIN)
    r = client.post(
        f"{API}/golden/candidates/export",
        json={"doc_ids": ["D3"], "department": "생산제조"},
    )
    wb = load_workbook(io.BytesIO(r.content))
    doc_ids = [row[0].value for row in wb.active.iter_rows(min_row=2)]
    assert doc_ids == ["D3"]


def test_export_row_carries_decision_reason_as_opinion(pool, client):
    """의견 = 결정 사유(reason) 그대로. 새 입력칸 없이도 엑셀에 나온다."""
    svc = pgs.ProxyGoldCandidateService()
    svc.decide(doc_id="D1", action="change", actor_id="reviewer-01", grade="S2", reason="현장 문의 결과 S2 로 판단")
    _as(ADMIN)
    r = client.post(f"{API}/golden/candidates/export", json={"doc_ids": ["D1"]})
    wb = load_workbook(io.BytesIO(r.content))
    row = next(wb.active.iter_rows(min_row=2, values_only=True))
    # doc_id, department, info_type, status, proposed_grade, final_grade, reviewer, reason, decided_at
    assert row[0] == "D1"
    assert row[5] == "S2"            # final_grade
    assert row[6] == "reviewer-01"   # reviewer
    assert row[7] == "현장 문의 결과 S2 로 판단"


def test_export_is_admin_only(pool, client):
    _as(REVIEWER)
    r = client.post(f"{API}/golden/candidates/export", json={})
    assert r.status_code == 403
